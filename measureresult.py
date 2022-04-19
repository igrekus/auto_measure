import os
import random

import openpyxl
import pandas as pd

from collections import defaultdict
from openpyxl.chart import LineChart, Series, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.cell import Cell
from textwrap import dedent

from forgot_again.file import load_ast_if_exists, pprint_to_file, make_dirs, open_explorer_at
from forgot_again.string import now_timestamp

GIGA = 1_000_000_000
MEGA = 1_000_000
KILO = 1_000
MILLI = 1 / 1_000


class MeasureResult:
    device = 'vco'
    measurement_name = 'tune'
    path = 'xlsx'

    def __init__(self):
        self._secondaryParams = dict()
        self._raw = list()

        self._report = dict()

        self._processed = list()

        self.ready = False

        self.data1 = defaultdict(list)
        self.data2 = defaultdict(list)
        self.data3 = defaultdict(list)
        self.data4 = defaultdict(list)

        self._table_data = list()
        self._table_header = list()

    def __bool__(self):
        return self.ready

    def process(self):
        self._prepare_table_data()
        self.ready = True

    def _process_point(self, data):
        series1 = data['series1']
        series2 = data['series2']
        series3 = data['series3']
        series4 = data['series4']

        x1 = data['x1']
        x2 = data['x2']
        x3 = data['x3']
        x4 = data['x4']

        y1 = data['y1']
        y2 = data['y2']
        y3 = data['y3']
        y4 = data['y4']

        self.data1[series1].append([x1, y1])
        self.data2[series2].append([x2, y2])
        self.data3[series3].append([x3, y3])
        self.data4[series4].append([x4, y4])

        self._processed.append({**data})

    def clear(self):
        self._secondaryParams.clear()
        self._raw.clear()

        self._report.clear()

        self._processed.clear()

        self.data1.clear()
        self.data2.clear()
        self.data3.clear()
        self.data4.clear()

        # self._table_data.clear()
        self._table_header.clear()

        self.ready = False

    def set_secondary_params(self, params):
        self._secondaryParams = dict(**params.params)

    def add_point(self, data):
        self._raw.append(data)
        self._process_point(data)

    @property
    def report(self):
        return dedent("""report""")

    def _prepare_table_data(self):
        table_file = './tables/stat_table.xlsx'

        if not os.path.isfile(table_file):
            return

        wb = openpyxl.load_workbook(table_file)
        ws = wb.active

        rows = list(ws.rows)
        self._table_header = [row.value for row in rows[0][1:]]

        gens = [
            [rows[1][j].value, rows[2][j].value, rows[3][j].value]
            for j in range(1, ws.max_column)
        ]

        self._table_data.append([self._gen_value(col) for col in gens])

    def _gen_value(self, data):
        if not data:
            return '-'
        if '-' in data:
            return '-'
        span, step, mean = data
        start = mean - span
        stop = mean + span
        if span == 0 or step == 0:
            return mean
        return round(random.randint(0, int((stop - start) / step)) * step + start, 2)

    def get_result_table_data(self):
        print(self._table_header)
        print(self._table_data)
        return list(self._table_header), list(self._table_data)


def _add_chart(ws, xs, ys, title, loc, curve_labels=None, ax_titles=None):
    chart = LineChart()

    for y, label in zip(ys, curve_labels):
        ser = Series(y, title=label)
        chart.append(ser)

    chart.set_categories(xs)
    chart.title = title

    chart.x_axis.minorGridlines = ChartLines()
    chart.x_axis.tickLblPos = 'low'

    if ax_titles:
        chart.x_axis.title = ax_titles[0]
        chart.y_axis.title = ax_titles[1]
    # chart.x_axis.tickLblSkip = 3

    ws.add_chart(chart, loc)


def _find_deltas(harm, origin):
    return [[main['u_control'], -(main['p_out'] - harm[1])] for harm, main in zip(harm, origin)]
